from fastapi import FastAPI, UploadFile, File, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
import sqlite3, io, re, openpyxl, os, requests as _http
from datetime import datetime, date as date_type
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = FastAPI(title="IT Inventory API")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

# ── Database setup ─────────────────────────────────────────────────────────
# If TURSO_URL + TURSO_TOKEN are set → use Turso via HTTP (no Rust needed)
# Otherwise → use local SQLite file (local development)
TURSO_URL   = os.environ.get("TURSO_URL", "")
TURSO_TOKEN = os.environ.get("TURSO_TOKEN", "")
USE_TURSO   = bool(TURSO_URL and TURSO_TOKEN)

DATA_DIR = os.environ.get("DATA_DIR", os.path.dirname(os.path.abspath(__file__)))
os.makedirs(DATA_DIR, exist_ok=True)
DB_PATH = os.path.join(DATA_DIR, "inventory.db")

# ── Pure-Python Turso HTTP client (no Rust / no libsql-experimental) ──────
def _tv(p):
    """Encode a Python value into a Turso typed-value dict."""
    if p is None:                        return {"type":"null","value":None}
    if isinstance(p, bool):              return {"type":"integer","value":str(int(p))}
    if isinstance(p, int):               return {"type":"integer","value":str(p)}
    if isinstance(p, float):             return {"type":"float","value":str(p)}
    return {"type":"text","value":str(p)}

def _dv(v):
    """Decode a Turso typed-value dict back to a Python value."""
    if v is None: return None
    t, raw = v.get("type",""), v.get("value")
    if t == "null" or raw is None: return None
    if t == "integer": return int(raw)
    if t == "float":   return float(raw)
    return raw

class _TursoRow:
    """sqlite3.Row-compatible object: access by name or index."""
    __slots__ = ("_c","_v","_m")
    def __init__(self, cols, vals):
        object.__setattr__(self,"_c",cols)
        object.__setattr__(self,"_v",vals)
        object.__setattr__(self,"_m",dict(zip(cols,vals)))
    def __getitem__(self,k): return self._v[k] if isinstance(k,int) else self._m[k]
    def keys(self): return list(self._c)
    def __iter__(self): return iter(self._v)
    def __len__(self): return len(self._v)

class _TursoCursor:
    def __init__(self, cols, raw_rows, as_row):
        mk = (lambda r: _TursoRow(cols,[_dv(v) for v in r])) if as_row \
             else (lambda r: tuple(_dv(v) for v in r))
        self._rows = [mk(r) for r in raw_rows]
        self._i = 0
    def fetchall(self):  return self._rows
    def fetchone(self):
        r = self._rows[self._i] if self._i < len(self._rows) else None
        self._i += 1; return r
    def __iter__(self): return iter(self._rows)

class _TursoConn:
    """HTTP connection to Turso that looks like sqlite3.Connection."""
    def __init__(self, url, token):
        base = url.replace("libsql://","https://")
        self._ep   = base.rstrip("/") + "/v2/pipeline"
        self._hdrs = {"Authorization":f"Bearer {token}","Content-Type":"application/json"}
        self.row_factory = None
    def execute(self, sql, parameters=None):
        stmt = {"sql": sql}
        if parameters: stmt["args"] = [_tv(p) for p in parameters]
        r = _http.post(self._ep, headers=self._hdrs,
                       json={"requests":[{"type":"execute","stmt":stmt},{"type":"close"}]},
                       timeout=30)
        r.raise_for_status()
        res = r.json()["results"][0]
        if res.get("type") == "error":
            raise Exception(res.get("error",{}).get("message","Turso error"))
        result = res["response"]["result"]
        cols = [c["name"] for c in result.get("cols",[])]
        return _TursoCursor(cols, result.get("rows",[]), self.row_factory == sqlite3.Row)
    def commit(self): pass   # Turso auto-commits each statement
    def close(self):  pass

# ── get_db ─────────────────────────────────────────────────────────────────
def get_db():
    if USE_TURSO:
        conn = _TursoConn(TURSO_URL, TURSO_TOKEN)
        conn.row_factory = sqlite3.Row
        return conn
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def header_to_field(header: str) -> str:
    """'SYSTEM S.NO' → 'system_s_no',  'SR. NO.' → 'sr_no'"""
    s = header.strip().lower()
    s = re.sub(r'[^a-z0-9]+', '_', s)
    s = re.sub(r'_+', '_', s)
    return s.strip('_') or 'col'

def cell_to_str(val) -> str:
    """Convert any Excel cell value to a clean string, handling dates properly."""
    if val is None:
        return ""
    # datetime / date objects (openpyxl auto-converts formatted date cells)
    if isinstance(val, datetime):
        return val.strftime("%d/%m/%Y")
    if isinstance(val, date_type):
        return val.strftime("%d/%m/%Y")
    s = str(val).strip()
    # Excel serial date numbers (e.g. 41255 = 12/12/2012)
    if s.isdigit() and 30000 < int(s) < 60000:
        try:
            from datetime import timedelta
            d = datetime(1899, 12, 30) + timedelta(days=int(s))
            return d.strftime("%d/%m/%Y")
        except Exception:
            pass
    return s

def get_col_meta(conn):
    """Return [{field_name, excel_header}] ordered by display_order."""
    rows = conn.execute(
        "SELECT field_name, excel_header FROM col_meta ORDER BY display_order"
    ).fetchall()
    return [dict(r) for r in rows]

# ── DB init ───────────────────────────────────────────────────────────────
def init_db():
    conn = get_db()
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS systems (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            created_at TEXT DEFAULT (datetime('now','localtime')),
            updated_at TEXT DEFAULT (datetime('now','localtime'))
        );
        CREATE TABLE IF NOT EXISTS col_meta (
            field_name    TEXT PRIMARY KEY,
            excel_header  TEXT,
            display_order INTEGER DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS history (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            system_id   INTEGER,
            field_name  TEXT,
            old_value   TEXT,
            new_value   TEXT,
            changed_at  TEXT DEFAULT (datetime('now','localtime')),
            FOREIGN KEY (system_id) REFERENCES systems(id)
        );
    """)
    conn.commit()

    # Migrate old fixed-schema DBs: populate col_meta from existing columns
    if conn.execute("SELECT COUNT(*) FROM col_meta").fetchone()[0] == 0:
        old_display = {
            "sr_no":           "SR. NO.",
            "room_no":         "ROOM NO.",
            "section":         "SECTION",
            "system_make":     "SYSTEM MAKE",
            "model":           "MODEL",
            "system_sno":      "SYSTEM S.NO",
            "tft_no":          "TFT NO",
            "date_procurement":"DATE OF PROCUREMENT",
            "procurement_amt": "ORIGINAL PROCUREMENT COST",
            "operating_system":"OPERATING SYSTEM",
            "processor":       "PROCESSOR",
            "ram":             "RAM",
            "storage":         "STORAGE",
            "ups":             "UPS",
            "computer_name":   "COMPUTER NAME",
            "printer_name":    "PRINTER NAME",
            "printer_sr_no":   "PRINTER SR. NO.",
            "timestamp":       "TIMESTAMP",
            "source_sheet":    "__SOURCE_SHEET__",
        }
        skip = {"id", "created_at", "updated_at"}
        existing_cols = conn.execute("PRAGMA table_info(systems)").fetchall()
        for order, col in enumerate(existing_cols):
            fname = col[1]
            if fname not in skip:
                header = old_display.get(fname, fname.upper().replace("_", " "))
                conn.execute(
                    "INSERT OR IGNORE INTO col_meta (field_name, excel_header, display_order) VALUES (?,?,?)",
                    [fname, header, order]
                )
        conn.commit()

    conn.close()

init_db()

# ── Models ────────────────────────────────────────────────────────────────
class SystemUpdate(BaseModel):
    field: str
    value: str

class BulkDeleteRequest(BaseModel):
    ids: list

# ── Routes ────────────────────────────────────────────────────────────────

@app.get("/")
def root():
    return {"message": "IT Inventory API is running!"}

@app.get("/columns")
def get_columns():
    """Return [{field_name, excel_header}] ordered by display_order."""
    conn = get_db()
    cols = get_col_meta(conn)
    conn.close()
    return cols

@app.post("/sheets")
async def get_sheets(file: UploadFile = File(...)):
    """Return all sheet names and accurate data-row counts from an uploaded Excel file."""
    contents = await file.read()
    try:
        # Do NOT use read_only=True — it can silently fail on some Excel files
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    except Exception as e:
        raise HTTPException(400, f"Cannot read file: {e}. Make sure it is a valid .xlsx file.")
    sheets = []
    for name in wb.sheetnames:
        try:
            ws = wb[name]
            # Use the same smart header detection as upload:
            # find the first row with 3+ non-empty cells → that is the header row.
            # Everything after is data.
            header_row = None
            for row_idx in range(1, 6):
                cells = [c for c in ws[row_idx] if c.value not in (None, "")]
                if len(cells) >= 3:
                    header_row = row_idx
                    break
            if header_row is None:
                # Sheet has no proper header → skip / show as 0 rows
                sheets.append({"name": name, "rows": 0})
                continue
            data_rows = sum(
                1 for row in ws.iter_rows(min_row=header_row + 1, values_only=True)
                if any(v not in (None, "") for v in row)
            )
            sheets.append({"name": name, "rows": data_rows})
        except Exception:
            sheets.append({"name": name, "rows": 0})
    return {"sheets": sheets}

@app.post("/upload")
async def upload_excel(file: UploadFile = File(...), clear_existing: bool = False, sheet_name: str = ""):
    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    # Detect header row: scan rows 1-5, pick first row with 3+ non-empty cells.
    # This skips single-cell title rows like "4TH FLOOR" or "3RD Floor" that
    # appear before the real column headers in some sheets.
    raw_headers = {}
    data_start = 2
    for row_idx in range(1, 6):
        candidate = {}
        for cell in ws[row_idx]:
            if cell.value:
                val = str(cell.value).strip().replace("\n", " ")
                val = re.sub(r'\s+', ' ', val).strip()
                if val:
                    candidate[val] = cell.column - 1
        if len(candidate) >= 3:
            raw_headers = candidate
            data_start = row_idx + 1
            break

    if not raw_headers:
        raise HTTPException(400, "Could not detect header row in Excel file")

    conn = get_db()

    if clear_existing:
        conn.execute("DELETE FROM history")
        conn.execute("DELETE FROM systems")
        conn.execute("DELETE FROM col_meta")
        conn.commit()

    # Existing DB columns and col_meta (by header for dedup)
    existing_db_cols = {row[1] for row in conn.execute("PRAGMA table_info(systems)").fetchall()}
    meta_by_header   = {
        row["excel_header"].upper(): row["field_name"]
        for row in get_col_meta(conn)
    }

    col_map = {}   # field_name → Excel col index
    for order, (header, col_idx) in enumerate(raw_headers.items()):
        # Reuse existing field_name if same header already in col_meta
        existing = meta_by_header.get(header.upper())
        field_name = existing if existing else header_to_field(header)

        if not field_name or field_name in {"id", "created_at", "updated_at"}:
            continue

        # Add DB column if new
        if field_name not in existing_db_cols:
            conn.execute(f'ALTER TABLE systems ADD COLUMN "{field_name}" TEXT DEFAULT ""')
            conn.commit()
            existing_db_cols.add(field_name)

        # Upsert col_meta
        conn.execute("""
            INSERT INTO col_meta (field_name, excel_header, display_order) VALUES (?,?,?)
            ON CONFLICT(field_name) DO UPDATE SET
                excel_header  = excluded.excel_header,
                display_order = excluded.display_order
        """, [field_name, header, order])

        col_map[field_name] = col_idx

    conn.commit()

    # Ensure source_sheet column exists (not shown in UI columns, handled separately)
    if "source_sheet" not in existing_db_cols:
        conn.execute('ALTER TABLE systems ADD COLUMN "source_sheet" TEXT DEFAULT ""')
        conn.commit()

    # Import rows
    DASH_PLACEHOLDERS = {"—", "–", "−", "﹣", "―", "-", "N/A", "NA", "n/a", "nil", "NULL", "null"}
    imported = 0
    fields = list(col_map.keys())

    for row in ws.iter_rows(min_row=data_start, values_only=True):
        if not any(row):
            continue
        vals = []
        for field in fields:
            ci = col_map[field]
            raw = row[ci] if ci < len(row) else None
            val = cell_to_str(raw)
            if val in DASH_PLACEHOLDERS:
                val = ""
            vals.append(val)

        all_fields = fields + ["source_sheet"]
        all_vals   = vals + [ws.title]
        conn.execute(
            f'INSERT INTO systems ({",".join(chr(34)+f+chr(34) for f in all_fields)}) '
            f'VALUES ({",".join(["?"]*len(all_fields))})',
            all_vals
        )
        imported += 1

    conn.commit()
    conn.close()
    return {"message": f"Imported {imported} records successfully"}


@app.get("/filter-options")
def get_filter_options():
    conn = get_db()
    meta = get_col_meta(conn)
    db_fields = [m["field_name"] for m in meta]

    section_field = next((f for f in db_fields if "section" in f), None)
    room_field    = next((f for f in db_fields if "room" in f), None)

    sections, rooms = [], []
    if section_field:
        sections = [r[0] for r in conn.execute(
            f'SELECT DISTINCT "{section_field}" FROM systems WHERE "{section_field}" != "" ORDER BY "{section_field}"'
        ).fetchall()]
    if room_field:
        rooms = [r[0] for r in conn.execute(
            f'SELECT DISTINCT "{room_field}" FROM systems WHERE "{room_field}" != "" ORDER BY "{room_field}"'
        ).fetchall()]

    # Sheet quick-switch: list all unique source sheets
    db_cols = {row[1] for row in conn.execute("PRAGMA table_info(systems)").fetchall()}
    sheets = []
    if "source_sheet" in db_cols:
        sheets = [r[0] for r in conn.execute(
            'SELECT DISTINCT source_sheet FROM systems WHERE source_sheet != "" ORDER BY source_sheet'
        ).fetchall()]

    conn.close()
    return {"sections": sections, "rooms": rooms, "sheets": sheets,
            "section_field": section_field, "room_field": room_field}


@app.get("/active-columns")
def get_active_columns(sheet: str = ""):
    """Return field_names that have at least one non-empty value for the given sheet."""
    conn = get_db()
    meta = get_col_meta(conn)
    if not sheet:
        conn.close()
        return [m["field_name"] for m in meta]
    active = []
    for m in meta:
        fn = m["field_name"]
        row = conn.execute(
            f'SELECT 1 FROM systems WHERE source_sheet=? AND "{fn}"!="" LIMIT 1', [sheet]
        ).fetchone()
        if row:
            active.append(fn)
    conn.close()
    return active


@app.get("/systems")
def get_systems(
    search: str = "", page: int = 1, limit: int = 20,
    sort_by: str = "", sort_dir: str = "asc",
    section: str = "", room_no: str = "", sheet: str = ""
):
    conn = get_db()
    meta      = get_col_meta(conn)
    db_fields = [m["field_name"] for m in meta]
    offset    = (page - 1) * limit
    order     = "DESC" if sort_dir == "desc" else "ASC"

    # Default sort: first sr/no-like field
    if not sort_by or sort_by not in db_fields + ["id"]:
        sort_by = next((f for f in db_fields if "sr" in f), db_fields[0] if db_fields else "id")

    # Numeric sort for serial/number fields
    numeric_hint = any(k in sort_by for k in ["sr", "no", "num", "count"])
    sort_expr = f'CAST("{sort_by}" AS INTEGER)' if numeric_hint else f'"{sort_by}"'

    section_field = next((f for f in db_fields if "section" in f), None)
    room_field    = next((f for f in db_fields if "room" in f), None)

    conditions, params = [], []
    if search and db_fields:
        like = f"%{search}%"
        conditions.append("(" + " OR ".join(f'"{f}" LIKE ?' for f in db_fields) + ")")
        params.extend([like] * len(db_fields))
    if section and section_field:
        conditions.append(f'"{section_field}" = ?')
        params.append(section)
    if room_no and room_field:
        conditions.append(f'"{room_field}" = ?')
        params.append(room_no)
    if sheet:
        conditions.append('"source_sheet" = ?')
        params.append(sheet)

    where = f"WHERE {' AND '.join(conditions)}" if conditions else ""

    rows  = conn.execute(
        f'SELECT * FROM systems {where} ORDER BY {sort_expr} {order} LIMIT ? OFFSET ?',
        params + [limit, offset]
    ).fetchall()
    total = conn.execute(f"SELECT COUNT(*) FROM systems {where}", params).fetchone()[0]

    conn.close()
    return {"data": [dict(r) for r in rows], "total": total, "page": page, "limit": limit}


@app.get("/systems/{system_id}")
def get_system(system_id: int):
    conn = get_db()
    row = conn.execute("SELECT * FROM systems WHERE id=?", [system_id]).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "System not found")
    return dict(row)


@app.patch("/systems/{system_id}")
def update_system(system_id: int, update: SystemUpdate):
    conn = get_db()
    db_fields = [m["field_name"] for m in get_col_meta(conn)]
    if update.field not in db_fields:
        conn.close()
        raise HTTPException(400, f"Invalid field: {update.field}")

    row = conn.execute("SELECT * FROM systems WHERE id=?", [system_id]).fetchone()
    if not row:
        conn.close()
        raise HTTPException(404, "System not found")

    old_val = dict(row).get(update.field, "")
    conn.execute(
        f'UPDATE systems SET "{update.field}"=?, updated_at=datetime("now","localtime") WHERE id=?',
        [update.value, system_id]
    )
    conn.execute(
        "INSERT INTO history (system_id, field_name, old_value, new_value) VALUES (?,?,?,?)",
        [system_id, update.field, old_val, update.value]
    )
    conn.commit()
    conn.close()
    return {"message": "Updated", "field": update.field, "old": old_val, "new": update.value}


@app.get("/systems/{system_id}/history")
def get_history(system_id: int):
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM history WHERE system_id=? ORDER BY changed_at DESC", [system_id]
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


@app.post("/systems")
async def create_system(request: Request):
    data = await request.json()
    conn = get_db()
    db_fields = [m["field_name"] for m in get_col_meta(conn)]

    fields = [f for f in db_fields if f in data and str(data[f]).strip()]
    if not fields:
        conn.close()
        raise HTTPException(400, "No valid fields provided")

    vals = [str(data[f]).strip() for f in fields]
    conn.execute(
        f'INSERT INTO systems ({",".join(chr(34)+f+chr(34) for f in fields)}) '
        f'VALUES ({",".join(["?"]*len(fields))})',
        vals
    )
    conn.commit()
    conn.close()
    return {"message": "System added successfully"}


@app.post("/systems/bulk-delete")
def bulk_delete_systems(req: BulkDeleteRequest):
    conn = get_db()
    for sid in req.ids:
        conn.execute("DELETE FROM systems WHERE id=?", [sid])
        conn.execute("DELETE FROM history WHERE system_id=?", [sid])
    conn.commit()
    conn.close()
    return {"message": f"Deleted {len(req.ids)} systems"}


@app.delete("/systems/all")
def delete_all_systems():
    conn = get_db()
    conn.execute("DELETE FROM history")
    conn.execute("DELETE FROM systems")
    conn.execute("DELETE FROM col_meta")
    conn.commit()
    conn.close()
    return {"message": "All data cleared"}


@app.delete("/systems/{system_id}")
def delete_system(system_id: int):
    conn = get_db()
    conn.execute("DELETE FROM systems WHERE id=?", [system_id])
    conn.execute("DELETE FROM history WHERE system_id=?", [system_id])
    conn.commit()
    conn.close()
    return {"message": "Deleted"}


@app.get("/download")
def download_excel():
    conn = get_db()
    meta = get_col_meta(conn)
    rows = conn.execute("SELECT * FROM systems ORDER BY id").fetchall()
    conn.close()

    if not meta:
        raise HTTPException(400, "No column metadata. Import an Excel file first.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "System Inventory"

    NAVY, BLUE, LTBLUE = "1F3864", "2E75B6", "DCE6F1"
    thin   = Side(style="thin", color="B8CCE4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    col_names   = [m["excel_header"] for m in meta]
    field_names = [m["field_name"]   for m in meta]
    num_cols    = len(col_names)

    # Title row
    ws.merge_cells(f"A1:{get_column_letter(num_cols)}1")
    t = ws["A1"]
    t.value     = "SYSTEM INVENTORY REGISTER"
    t.font      = Font(name="Calibri", bold=True, size=16, color="FFFFFF")
    t.fill      = PatternFill("solid", start_color=BLUE)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Header row
    for ci, h in enumerate(col_names, 1):
        c = ws.cell(row=2, column=ci, value=h)
        c.font      = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
        c.fill      = PatternFill("solid", start_color=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border    = border
    ws.row_dimensions[2].height = 40

    # Data rows
    for ri, row in enumerate(rows, 3):
        d     = dict(row)
        shade = LTBLUE if ri % 2 == 0 else "FFFFFF"
        for ci, fname in enumerate(field_names, 1):
            c = ws.cell(row=ri, column=ci, value=d.get(fname, ""))
            c.font      = Font(name="Calibri", size=10)
            c.fill      = PatternFill("solid", start_color=shade)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border    = border

    # Auto column widths
    for i, fname in enumerate(field_names, 1):
        best = max(len(str(ws.cell(row=r, column=i).value or "")) for r in range(2, min(ws.max_row+1, 12)))
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(best + 4, 40))

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(num_cols)}2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname = f"SystemInventory_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"}
    )


@app.get("/stats")
def get_stats():
    conn   = get_db()
    total  = conn.execute("SELECT COUNT(*) FROM systems").fetchone()[0]
    meta   = get_col_meta(conn)
    fields = [m["field_name"] for m in meta]

    section_field = next((f for f in fields if "section" in f), None)
    make_field    = next((f for f in fields if "make" in f or "brand" in f), None)

    sections = []
    if section_field:
        sections = [dict(r) for r in conn.execute(
            f'SELECT "{section_field}" as section, COUNT(*) as cnt FROM systems '
            f'GROUP BY "{section_field}" ORDER BY cnt DESC'
        ).fetchall()]

    makes = []
    if make_field:
        makes = [dict(r) for r in conn.execute(
            f'SELECT "{make_field}" as system_make, COUNT(*) as cnt FROM systems '
            f'GROUP BY "{make_field}" ORDER BY cnt DESC'
        ).fetchall()]

    recent = [dict(r) for r in conn.execute(
        "SELECT * FROM history ORDER BY changed_at DESC LIMIT 10"
    ).fetchall()]

    conn.close()
    return {"total": total, "by_section": sections, "by_make": makes, "recent_changes": recent}



